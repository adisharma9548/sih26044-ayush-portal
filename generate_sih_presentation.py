import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_sih_presentation(output_path="SIH26044_NodalConnector_Idea_Presentation.pptx"):
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors (SIH 2026 Official Palette)
    NAVY_PRIMARY = RGBColor(26, 54, 93)      # #1A365D - Deep SIH Navy
    ORANGE_ACCENT = RGBColor(217, 119, 6)    # #D97706 - SIH Amber / Orange
    BLUE_BANNER = RGBColor(0, 112, 186)      # #0070BA - SIH Light Blue Bar
    DARK_TEXT = RGBColor(30, 41, 59)         # #1E293B - Slate Charcoal
    MUTED_TEXT = RGBColor(100, 116, 139)     # #64748B - Slate Gray
    CARD_BG = RGBColor(248, 250, 252)        # #F8FAFC - Clean Card Background
    CARD_BORDER = RGBColor(226, 232, 240)    # #E2E8F0 - Subtle Border
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BLUE_BG = RGBColor(238, 246, 255)  # Light Blue Tint
    SS_BOX_BG = RGBColor(241, 245, 249)      # Screenshot placeholder background
    SS_BOX_BORDER = RGBColor(14, 116, 144)   # Cyan / Teal Border for screenshot placeholders

    def add_bottom_bar(slide, slide_num):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.0), Inches(13.333), Inches(0.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = BLUE_BANNER
        bar.line.color.rgb = BLUE_BANNER

        tx_box = slide.shapes.add_textbox(Inches(3.0), Inches(7.05), Inches(7.333), Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "@SIH Idea submission- Template"
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Calibri"
        p.font.size = Pt(11)
        p.font.color.rgb = WHITE

        num_box = slide.shapes.add_textbox(Inches(12.2), Inches(7.05), Inches(0.9), Inches(0.4))
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = str(slide_num)
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.font.name = "Calibri"
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = WHITE

    def add_sih_header_badge(slide):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.1), Inches(0.2), Inches(1.9), Inches(0.8))
        badge.fill.solid()
        badge.fill.fore_color.rgb = WHITE
        badge.line.color.rgb = CARD_BORDER
        badge.line.width = Pt(1)

        tf = badge.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p1 = tf.paragraphs[0]
        p1.text = "SMART INDIA"
        p1.alignment = PP_ALIGN.CENTER
        p1.font.name = "Arial"
        p1.font.bold = True
        p1.font.size = Pt(9)
        p1.font.color.rgb = NAVY_PRIMARY

        p2 = tf.add_paragraph()
        p2.text = "HACKATHON 2026"
        p2.alignment = PP_ALIGN.CENTER
        p2.font.name = "Arial"
        p2.font.bold = True
        p2.font.size = Pt(8)
        p2.font.color.rgb = ORANGE_ACCENT

    def add_screenshot_placeholder(slide, x, y, w, h, title, route_url, description):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        box.fill.solid()
        box.fill.fore_color.rgb = SS_BOX_BG
        box.line.color.rgb = SS_BOX_BORDER
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.15)
        tf.margin_bottom = Inches(0.15)

        p1 = tf.paragraphs[0]
        p1.text = f"📷 [REAL SCREENSHOT: {title}]"
        p1.alignment = PP_ALIGN.CENTER
        p1.font.name = "Arial"
        p1.font.bold = True
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = NAVY_PRIMARY
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = f"Route URL: {route_url}"
        p2.alignment = PP_ALIGN.CENTER
        p2.font.name = "Consolas"
        p2.font.bold = True
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = ORANGE_ACCENT
        p2.space_after = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = description
        p3.alignment = PP_ALIGN.CENTER
        p3.font.name = "Calibri"
        p3.font.size = Pt(9)
        p3.font.color.rgb = DARK_TEXT

    # ==========================================
    # SLIDE 1: TITLE SLIDE (MVP FOCUS)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)

    title_box = s1.shapes.add_textbox(Inches(1.0), Inches(0.4), Inches(11.333), Inches(1.1))
    tf1 = title_box.text_frame
    p_title = tf1.paragraphs[0]
    p_title.text = "SMART INDIA HACKATHON 2026"
    p_title.alignment = PP_ALIGN.CENTER
    p_title.font.name = "Georgia"
    p_title.font.bold = True
    p_title.font.size = Pt(36)
    p_title.font.color.rgb = NAVY_PRIMARY

    p_sub = tf1.add_paragraph()
    p_sub.text = "NODAL CONNECTOR (WORKING MVP)"
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.font.name = "Arial"
    p_sub.font.bold = True
    p_sub.font.size = Pt(22)
    p_sub.font.color.rgb = ORANGE_ACCENT

    # Left content card with problem statement details
    details_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(8.3), Inches(5.15))
    details_card.fill.solid()
    details_card.fill.fore_color.rgb = WHITE
    details_card.line.color.rgb = CARD_BORDER
    details_card.line.width = Pt(1.5)

    tf_details = details_card.text_frame
    tf_details.word_wrap = True
    tf_details.margin_left = Inches(0.4)
    tf_details.margin_top = Inches(0.35)
    tf_details.margin_right = Inches(0.4)

    fields = [
        ("❑ Problem Statement ID : ", "SIH26044"),
        ("❑ Problem Statement Title : ", "Portal for Academia - Industry collaboration for Skill Mapping, Internships and Placement"),
        ("❑ Theme : ", "Smart Automation"),
        ("❑ PS Category : ", "Software"),
        ("❑ Project Status : ", "Functional MVP Prototype (Full-Stack Built & Deployed)"),
        ("❑ Team ID : ", "[Insert Team ID]"),
        ("❑ Team Name (Registered on portal) : ", "[Insert Registered Team Name]")
    ]

    for i, (label, val) in enumerate(fields):
        p = tf_details.paragraphs[0] if i == 0 else tf_details.add_paragraph()
        p.space_after = Pt(10)
        r1 = p.add_run()
        r1.text = label
        r1.font.bold = True
        r1.font.name = "Calibri"
        r1.font.size = Pt(13.5)
        r1.font.color.rgb = NAVY_PRIMARY

        r2 = p.add_run()
        r2.text = val
        r2.font.bold = (i in [0, 2, 3, 4])
        r2.font.name = "Calibri"
        r2.font.size = Pt(13.5)
        r2.font.color.rgb = DARK_TEXT

    # Right Side MVP Snapshot Card
    right_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.4), Inches(1.85), Inches(3.2), Inches(5.15))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = LIGHT_BLUE_BG
    right_card.line.color.rgb = BLUE_BANNER
    right_card.line.width = Pt(1.2)

    tf_right = right_card.text_frame
    tf_right.word_wrap = True
    tf_right.margin_left = Inches(0.25)
    tf_right.margin_right = Inches(0.25)
    tf_right.margin_top = Inches(0.25)

    p_rt = tf_right.paragraphs[0]
    p_rt.text = "WORKING MVP HIGHLIGHTS"
    p_rt.alignment = PP_ALIGN.CENTER
    p_rt.font.bold = True
    p_rt.font.size = Pt(13)
    p_rt.font.color.rgb = NAVY_PRIMARY

    spec_points = [
        "🌐 Live Web Portal:\nDeployed frontend & backend API",
        "🎯 4 Functional Roles:\nStudent, Industry, Faculty, Admin",
        "🧠 AI Diagnostic Tests:\nDynamic skill testing (Groq LLaMA 3.3)",
        "📹 Live WebRTC Rooms:\nZero-install video & collaborative whiteboard",
        "📊 Dynamic Skill Match:\nTransparent score based on real abilities"
    ]

    for pt in spec_points:
        p_pt = tf_right.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.space_after = Pt(8)
        p_pt.font.size = Pt(10)
        p_pt.font.color.rgb = DARK_TEXT

    add_sih_header_badge(s1)

    # ==========================================
    # SLIDE 2: NODAL CONNECTOR (PROPOSED SOLUTION)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bottom_bar(s2, 2)
    add_sih_header_badge(s2)

    t2 = s2.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10.0), Inches(0.6))
    tf2 = t2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "NODAL CONNECTOR"
    p2.font.name = "Georgia"
    p2.font.bold = True
    p2.font.size = Pt(28)
    p2.font.color.rgb = NAVY_PRIMARY

    sub2 = s2.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(6.0), Inches(0.4))
    tf_sub2 = sub2.text_frame
    p_sub2 = tf_sub2.paragraphs[0]
    p_sub2.text = "❖ Proposed Solution"
    p_sub2.font.name = "Arial"
    p_sub2.font.bold = True
    p_sub2.font.size = Pt(18)
    p_sub2.font.color.rgb = NAVY_PRIMARY

    # Left: 2 Columns of Core Pillars (Inches 0.8 to 8.2)
    col_w = Inches(3.6)
    card_y = Inches(1.45)
    card_h = Inches(5.2)

    pillars = [
        {
            "num": "01",
            "title": "Smart Skill Diagnostic & Gap Remediation",
            "items": [
                "Objective AI Diagnostic Assessment (Groq LLaMA 3.3) assessing real student competencies with authentic 0-100% scoring.",
                "UGC Section 22 Degree Normalizer with 500ms debounce matching engineering, bio-pharma, and healthcare degrees.",
                "Automated Skill Gap Matrix linking identified weaknesses directly to curated roadmap.sh learning pathways.",
                "Real Cohort Percentile Benchmarks calculated dynamically from peer test attempts without artificial score boosts."
            ]
        },
        {
            "num": "02",
            "title": "Unified Opportunities & Direct In-App Hiring",
            "items": [
                "Centralized Marketplace for verified experiential internships and entry-level jobs with stipends.",
                "Transparent Compatibility Score based on verified skills matching specific posting requirements.",
                "Real-Time Application Lifecycle: Applied → Under Review → Interview Scheduled → Offered / Rejected.",
                "Collaborative Industry Learning Programs: Companies sponsor structured modules for talent upskilling."
            ]
        }
    ]

    for idx, p_info in enumerate(pillars):
        x = Inches(0.8) + idx * (col_w + Inches(0.25))
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, card_y, col_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        header_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.12), card_y + Inches(0.12), col_w - Inches(0.24), Inches(0.75))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = NAVY_PRIMARY
        header_shape.line.fill.background()

        h_tf = header_shape.text_frame
        h_tf.word_wrap = True
        h_p1 = h_tf.paragraphs[0]
        h_p1.text = f"Pillar {p_info['num']}: {p_info['title']}"
        h_p1.font.name = "Arial"
        h_p1.font.bold = True
        h_p1.font.size = Pt(10.5)
        h_p1.font.color.rgb = WHITE
        h_p1.alignment = PP_ALIGN.CENTER

        content_box = s2.shapes.add_textbox(x + Inches(0.12), card_y + Inches(0.95), col_w - Inches(0.24), card_h - Inches(1.05))
        c_tf = content_box.text_frame
        c_tf.word_wrap = True
        for b_idx, bullet in enumerate(p_info["items"]):
            bp = c_tf.paragraphs[0] if b_idx == 0 else c_tf.add_paragraph()
            bp.text = "• " + bullet
            bp.space_after = Pt(8)
            bp.font.name = "Calibri"
            bp.font.size = Pt(10)
            bp.font.color.rgb = DARK_TEXT

    # Right: Real Visual / Screenshot Placement Box
    add_screenshot_placeholder(
        s2,
        Inches(8.35), Inches(1.45), Inches(4.2), Inches(5.2),
        "STUDENT DASHBOARD & AI ASSESSMENT",
        "/student/dashboard & /student/assessment",
        "Take a screenshot of the Student Dashboard showing the Skill Radar Chart, Active Applications, or the Diagnostic Test in action!"
    )

    # ==========================================
    # SLIDE 3: TECHNICAL APPROACH
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bottom_bar(s3, 3)
    add_sih_header_badge(s3)

    t3 = s3.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10.0), Inches(0.6))
    tf3 = t3.text_frame
    p3 = tf3.paragraphs[0]
    p3.text = "TECHNICAL APPROACH"
    p3.font.name = "Georgia"
    p3.font.bold = True
    p3.font.size = Pt(28)
    p3.font.color.rgb = NAVY_PRIMARY

    sub3 = s3.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(6.0), Inches(0.4))
    tf_sub3 = sub3.text_frame
    p_sub3 = tf_sub3.paragraphs[0]
    p_sub3.text = "❖ Technology Stack"
    p_sub3.font.name = "Arial"
    p_sub3.font.bold = True
    p_sub3.font.size = Pt(18)
    p_sub3.font.color.rgb = NAVY_PRIMARY

    # Left: 3 Tech Stack Domain Blocks
    tech_blocks = [
        {
            "layer": "Frontend (React 18 + Vite + Tailwind CSS)",
            "points": [
                "Single Page App (SPA) with Protected Routes for 4 personas.",
                "Zustand state store for user sessions & real-time alerts.",
                "Dynamic SVG skill radar charts & debounced UGC degree selector."
            ]
        },
        {
            "layer": "Backend & Real-Time (Node.js + Express + Socket.IO)",
            "points": [
                "Modular RESTful API with JWT authentication & role authorization.",
                "Socket.IO signaling for browser-to-browser WebRTC video/audio calls.",
                "Real-time synchronized whiteboard canvas for technical discussions."
            ]
        },
        {
            "layer": "AI & Database (Groq LLaMA 3.3 + MongoDB Atlas)",
            "points": [
                "Groq LPU high-speed inference for instant diagnostic question generation.",
                "17 Mongoose schemas with compound indexes for fast applicant search.",
                "roadmap.sh curriculum integration for automated student study roadmaps."
            ]
        }
    ]

    card_y_s3 = Inches(1.45)
    for idx, b in enumerate(tech_blocks):
        by = card_y_s3 + idx * Inches(1.72)
        b_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), by, Inches(6.8), Inches(1.58))
        b_card.fill.solid()
        b_card.fill.fore_color.rgb = CARD_BG
        b_card.line.color.rgb = CARD_BORDER
        b_card.line.width = Pt(1.5)

        btf = b_card.text_frame
        btf.word_wrap = True
        btf.margin_left = Inches(0.2)
        btf.margin_right = Inches(0.2)
        btf.margin_top = Inches(0.1)

        p_l = btf.paragraphs[0]
        p_l.text = b["layer"]
        p_l.font.name = "Arial"
        p_l.font.bold = True
        p_l.font.size = Pt(11.5)
        p_l.font.color.rgb = NAVY_PRIMARY

        for pt in b["points"]:
            p_pt = btf.add_paragraph()
            p_pt.text = "• " + pt
            p_pt.font.name = "Calibri"
            p_pt.font.size = Pt(9.5)
            p_pt.font.color.rgb = DARK_TEXT

    # Right: Real Visual / Screenshot Placement Box (Live WebRTC Meeting Room)
    add_screenshot_placeholder(
        s3,
        Inches(7.8), Inches(1.45), Inches(4.75), Inches(5.2),
        "LIVE WEBRTC INTERVIEW & WHITEBOARD",
        "/meetings (or direct interview room)",
        "Take a screenshot of the Live Meeting page showing the WebRTC video tiles, controls, and the interactive drawing canvas board!"
    )

    # ==========================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bottom_bar(s4, 4)
    add_sih_header_badge(s4)

    t4 = s4.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10.0), Inches(0.6))
    tf4 = t4.text_frame
    p4 = tf4.paragraphs[0]
    p4.text = "FEASIBILITY AND VIABILITY"
    p4.font.name = "Georgia"
    p4.font.bold = True
    p4.font.size = Pt(28)
    p4.font.color.rgb = NAVY_PRIMARY

    # Left Column: Feasibility, Viability, Challenges
    left_w = Inches(3.9)

    box_feas = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.15), left_w, Inches(1.65))
    box_feas.fill.solid()
    box_feas.fill.fore_color.rgb = CARD_BG
    box_feas.line.color.rgb = CARD_BORDER
    tf_f = box_feas.text_frame
    tf_f.margin_left = Inches(0.18)
    tf_f.margin_top = Inches(0.1)
    p = tf_f.paragraphs[0]
    p.text = "Feasibility"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = NAVY_PRIMARY
    f_items = [
        "Technical: Open web standards (React, Node, WebRTC) needing 0 external downloads.",
        "Operational: Lightweight UI runs on mobile devices and basic college Wi-Fi.",
        "Financial: Serverless cloud hosting keeps running costs minimal for the MVP."
    ]
    for it in f_items:
        p_it = tf_f.add_paragraph()
        p_it.text = "• " + it
        p_it.font.size = Pt(9)
        p_it.font.color.rgb = DARK_TEXT

    box_viab = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.95), left_w, Inches(1.65))
    box_viab.fill.solid()
    box_viab.fill.fore_color.rgb = CARD_BG
    box_viab.line.color.rgb = CARD_BORDER
    tf_v = box_viab.text_frame
    tf_v.margin_left = Inches(0.18)
    tf_v.margin_top = Inches(0.1)
    p = tf_v.paragraphs[0]
    p.text = "Viability"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = NAVY_PRIMARY
    v_items = [
        "High College Adoption: Helps placement cells connect students to real recruiters.",
        "Recruiter Incentive: Pre-assessed candidates with transparent skill matches.",
        "MVP Scalability: Clean modular code easily transitions into full production."
    ]
    for it in v_items:
        p_it = tf_v.add_paragraph()
        p_it.text = "• " + it
        p_it.font.size = Pt(9)
        p_it.font.color.rgb = DARK_TEXT

    box_chal = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.75), left_w, Inches(1.95))
    box_chal.fill.solid()
    box_chal.fill.fore_color.rgb = CARD_BG
    box_chal.line.color.rgb = CARD_BORDER
    tf_c = box_chal.text_frame
    tf_c.margin_left = Inches(0.18)
    tf_c.margin_top = Inches(0.1)
    p = tf_c.paragraphs[0]
    p.text = "Challenges"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = NAVY_PRIMARY
    c_items = [
        "Cold Start: Students having unverified or incomplete skill profiles.",
        "Curricular Lag: University syllabi updating slower than industry tech.",
        "Campus Firewalls: Restrictive college NAT networks blocking P2P video."
    ]
    for it in c_items:
        p_it = tf_c.add_paragraph()
        p_it.text = "• " + it
        p_it.font.size = Pt(9)
        p_it.font.color.rgb = DARK_TEXT

    # Center Column: Use Cases & Solutions
    mid_w = Inches(3.9)
    mid_x = Inches(4.9)

    box_uc = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mid_x, Inches(1.15), mid_w, Inches(2.65))
    box_uc.fill.solid()
    box_uc.fill.fore_color.rgb = CARD_BG
    box_uc.line.color.rgb = CARD_BORDER
    tf_uc = box_uc.text_frame
    tf_uc.margin_left = Inches(0.18)
    tf_uc.margin_top = Inches(0.1)
    p = tf_uc.paragraphs[0]
    p.text = "Use Cases"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = NAVY_PRIMARY
    uc_items = [
        "Student Career Discovery: Student identifies a skill gap in bio-pharma / coding, follows the roadmap, and applies for an internship.",
        "Recruiter Sourcing: Employer filters candidates by degree & verified skill compatibility, scheduling an instant video interview.",
        "Faculty Collaboration: Professor browses industry research immersion programs and schedules student workshops."
    ]
    for it in uc_items:
        p_it = tf_uc.add_paragraph()
        p_it.text = "• " + it
        p_it.font.size = Pt(9)
        p_it.font.color.rgb = DARK_TEXT

    box_sol = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mid_x, Inches(3.95), mid_w, Inches(2.75))
    box_sol.fill.solid()
    box_sol.fill.fore_color.rgb = LIGHT_BLUE_BG
    box_sol.line.color.rgb = BLUE_BANNER
    tf_s = box_sol.text_frame
    tf_s.margin_left = Inches(0.18)
    tf_s.margin_top = Inches(0.1)
    p = tf_s.paragraphs[0]
    p.text = "Solutions"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = NAVY_PRIMARY
    sol_items = [
        "AI Diagnostic Testing: Fast 10-question evaluation immediately verifies core skills.",
        "Industry Learning Tracks: Companies post custom learning modules to bridge syllabus lags.",
        "STUN/TURN Fallback: Integrated WebRTC fallback ensures video calls connect even on campus Wi-Fi."
    ]
    for it in sol_items:
        p_it = tf_s.add_paragraph()
        p_it.text = "• " + it
        p_it.font.size = Pt(9)
        p_it.font.color.rgb = DARK_TEXT

    # Right Column: Real Visual / Screenshot Placement Box
    add_screenshot_placeholder(
        s4,
        Inches(9.0), Inches(1.15), Inches(3.55), Inches(5.55),
        "INTERNSHIP & JOB MARKETPLACE",
        "/student/internships or /industry/post",
        "Take a screenshot of the Internships portal showing opportunity cards, stipend badges, and the 'Skill Match' indicator!"
    )

    # ==========================================
    # SLIDE 5: IMPACT AND BENEFITS
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bottom_bar(s5, 5)
    add_sih_header_badge(s5)

    t5 = s5.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10.0), Inches(0.6))
    tf5 = t5.text_frame
    p5 = tf5.paragraphs[0]
    p5.text = "IMPACT AND BENEFITS"
    p5.font.name = "Georgia"
    p5.font.bold = True
    p5.font.size = Pt(28)
    p5.font.color.rgb = NAVY_PRIMARY

    half_w = Inches(3.9)

    # Section 1: Benefits Of The Solution
    card_b = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.15), half_w, Inches(5.55))
    card_b.fill.solid()
    card_b.fill.fore_color.rgb = CARD_BG
    card_b.line.color.rgb = CARD_BORDER
    card_b.line.width = Pt(1.5)

    tf_b = card_b.text_frame
    tf_b.margin_left = Inches(0.2)
    tf_b.margin_right = Inches(0.2)
    tf_b.margin_top = Inches(0.18)
    p_bh = tf_b.paragraphs[0]
    p_bh.text = "❖ Benefits Of The Solution"
    p_bh.font.name = "Arial"
    p_bh.font.bold = True
    p_bh.font.size = Pt(14)
    p_bh.font.color.rgb = NAVY_PRIMARY
    p_bh.space_after = Pt(8)

    stakeholder_benefits = [
        ("For Students & Jobseekers", [
            "Authentic skill feedback with personalized learning roadmaps.",
            "Direct access to verified internship & job openings with stipends.",
            "Shareable digital portfolio of verified skills."
        ]),
        ("For Industry & Recruiters", [
            "Saves hours of screening through transparent skill-match scores.",
            "Schedule and conduct technical interviews directly in the portal.",
            "Sponsor tailored learning programs to groom talent."
        ]),
        ("For Academic Faculty", [
            "Real-time insight into in-demand industry competencies.",
            "Direct channel to connect students with corporate mentors."
        ])
    ]

    for title, points in stakeholder_benefits:
        p_st = tf_b.add_paragraph()
        p_st.text = title
        p_st.font.bold = True
        p_st.font.size = Pt(10.5)
        p_st.font.color.rgb = ORANGE_ACCENT
        p_st.space_after = Pt(2)
        for pt in points:
            p_p = tf_b.add_paragraph()
            p_p.text = "• " + pt
            p_p.font.size = Pt(9)
            p_p.font.color.rgb = DARK_TEXT
            p_p.space_after = Pt(2)

    # Section 2: Potential Impact On The Target Audience
    card_i = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.9), Inches(1.15), half_w, Inches(5.55))
    card_i.fill.solid()
    card_i.fill.fore_color.rgb = CARD_BG
    card_i.line.color.rgb = CARD_BORDER
    card_i.line.width = Pt(1.5)

    tf_i = card_i.text_frame
    tf_i.margin_left = Inches(0.2)
    tf_i.margin_right = Inches(0.2)
    tf_i.margin_top = Inches(0.18)
    p_ih = tf_i.paragraphs[0]
    p_ih.text = "❖ Potential Impact On The Target Audience"
    p_ih.font.name = "Arial"
    p_ih.font.bold = True
    p_ih.font.size = Pt(14)
    p_ih.font.color.rgb = NAVY_PRIMARY
    p_ih.space_after = Pt(8)

    impact_points = [
        ("Bridging the Employability Gap",
         "Helps students uncover what they don't know early in college, giving them actionable steps to become job-ready."),
        ("Democratizing Tier-2/3 College Access",
         "Gives students from smaller colleges direct visibility to companies without relying solely on campus placement drives."),
        ("Multi-Disciplinary Synergy",
         "Connects traditional sciences, Ayush bio-pharma, and computer engineering under one unified skill taxonomy."),
        ("Transparent Merit-Based Hiring",
         "Replaces inflated resumes with verified diagnostic results and real project showcases.")
    ]

    for title, desc in impact_points:
        p_imp = tf_i.add_paragraph()
        p_imp.text = title
        p_imp.font.bold = True
        p_imp.font.size = Pt(10.5)
        p_imp.font.color.rgb = NAVY_PRIMARY
        p_imp.space_after = Pt(1)

        p_desc = tf_i.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(9)
        p_desc.font.color.rgb = DARK_TEXT
        p_desc.space_after = Pt(4)

    # Right: Real Visual / Screenshot Placement Box
    add_screenshot_placeholder(
        s5,
        Inches(9.0), Inches(1.15), Inches(3.55), Inches(5.55),
        "RECRUITER APPLICANT MANAGEMENT",
        "/industry/applicants or /industry/candidates",
        "Take a screenshot of the Recruiter page showing candidate applications, compatibility score pills, and interview actions!"
    )

    # ==========================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bottom_bar(s6, 6)
    add_sih_header_badge(s6)

    t6 = s6.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10.0), Inches(0.6))
    tf6 = t6.text_frame
    p6 = tf6.paragraphs[0]
    p6.text = "RESEARCH AND REFERENCES"
    p6.font.name = "Georgia"
    p6.font.bold = True
    p6.font.size = Pt(28)
    p6.font.color.rgb = NAVY_PRIMARY

    sub6 = s6.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(8.0), Inches(0.4))
    tf_sub6 = sub6.text_frame
    p_sub6 = tf_sub6.paragraphs[0]
    p_sub6.text = "• Details / Links of the reference and research work"
    p_sub6.font.name = "Arial"
    p_sub6.font.bold = True
    p_sub6.font.size = Pt(18)
    p_sub6.font.color.rgb = NAVY_PRIMARY

    # Left: Citations Box
    ref_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(7.9), Inches(5.25))
    ref_card.fill.solid()
    ref_card.fill.fore_color.rgb = CARD_BG
    ref_card.line.color.rgb = CARD_BORDER
    ref_card.line.width = Pt(1.5)

    tf_ref = ref_card.text_frame
    tf_ref.margin_left = Inches(0.25)
    tf_ref.margin_right = Inches(0.25)
    tf_ref.margin_top = Inches(0.2)

    references = [
        ("National Education Policy (NEP 2020) — Ministry of Education",
         "Reference: NEP 2020 Chapter 11 on vocational training, industry linkage, and experiential learning credits.",
         "https://www.education.gov.in/sites/upload_files/mhrd/files/NEP_Final_English_0.pdf"),

        ("University Grants Commission (UGC) Degree Specification",
         "Reference: UGC Act 1956 Section 22 official nomenclature standardization framework, utilized in our degree selector.",
         "https://www.ugc.gov.in/pdfnews/spec_degrees.pdf"),

        ("AICTE National Internship Policy & Guidelines",
         "Reference: Framework for mandatory technical internships and industry-academia credit integration.",
         "https://internship.aicte-india.org/"),

        ("roadmap.sh Open-Source Curriculum Framework",
         "Reference: Community developer skill roadmaps used to curate learning paths for student remediation.",
         "https://github.com/kamranahmedse/developer-roadmap"),

        ("Working MVP Repository & Live Deployments (SIH26044)",
         "• Live Client: https://sih26044-ayush-portal.vercel.app\n• Live Backend API: https://sih26044-ayush-portal-production.up.railway.app\n• GitHub: https://github.com/adisharma9548/sih26044-ayush-portal",
         "Working prototype solving SIH26044 with full MongoDB Atlas persistence and live WebRTC signaling.")
    ]

    for i, (title, citation, link) in enumerate(references):
        p_title = tf_ref.paragraphs[0] if i == 0 else tf_ref.add_paragraph()
        p_title.text = f"{i+1}. {title}"
        p_title.font.bold = True
        p_title.font.size = Pt(10)
        p_title.font.color.rgb = NAVY_PRIMARY

        p_cit = tf_ref.add_paragraph()
        p_cit.text = citation
        p_cit.font.size = Pt(8.5)
        p_cit.font.color.rgb = DARK_TEXT

        p_lnk = tf_ref.add_paragraph()
        p_lnk.text = link
        p_lnk.font.size = Pt(8)
        p_lnk.font.italic = True
        p_lnk.font.color.rgb = BLUE_BANNER
        p_lnk.space_after = Pt(4)

    # Right: Real Visual / Screenshot Placement Box
    add_screenshot_placeholder(
        s6,
        Inches(8.9), Inches(1.45), Inches(3.65), Inches(5.25),
        "INSTITUTIONAL PROGRESS & ANALYTICS",
        "/admin/reports or /admin/progress",
        "Take a screenshot of the Institutional Admin analytics dashboard showing student placement trends and skill distributions!"
    )

    prs.save(output_path)
    print(f"Presentation saved successfully to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_sih_presentation()
